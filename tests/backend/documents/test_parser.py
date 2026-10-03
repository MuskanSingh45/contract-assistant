from pathlib import Path

import docx
import pymupdf
import pytest

from backend.documents.normalizer import split_long
from backend.documents.parser import (
    DOCX_MIME,
    PDF_MIME,
    TEXT_MIME,
    DocumentParseError,
    NoExtractableText,
    detect_mime_type,
    parse_document,
)

ROOT = Path(__file__).resolve().parents[3]
ACME_TXT = ROOT / "contracts/samples/simple/acme-services-agreement.txt"
ACME_PDF = ROOT / "contracts/samples/simple/acme-services-agreement.pdf"
ACME_DOCX = ROOT / "contracts/samples/simple/acme-services-agreement.docx"
RENEWAL_QUOTE = "at least ninety (90) days before the expiration"


def make_pdf(path: Path, pages: list[list[str]]) -> Path:
    """Each page is a list of paragraphs; each paragraph is inserted as its own text block."""
    doc = pymupdf.open()
    for paragraphs in pages:
        page = doc.new_page()
        y = 72
        for text in paragraphs:
            page.insert_textbox(pymupdf.Rect(72, y, 520, y + 80), text, fontsize=11)
            y += 100
    doc.save(path)
    doc.close()
    return path


def make_docx(path: Path, paragraphs: list[tuple[str, str]]) -> Path:
    document = docx.Document()
    for style, text in paragraphs:
        document.add_paragraph(text, style=style)
    document.save(path)
    return path


def by_text(segments, needle):
    return next(s for s in segments if needle in s.text)


def test_pdf_pages_and_numbered_heading_section(tmp_path):
    pdf = make_pdf(
        tmp_path / "two.pdf",
        [
            [
                "1. DEFINITIONS",
                "The Customer means Acme Corporation and its affiliates.",
            ],
            ["5 Payment Terms", "Invoices are payable within thirty days of receipt."],
        ],
    )
    segments, page_count = parse_document(str(pdf), PDF_MIME)
    assert page_count == 2
    assert [(s.seq, s.page, s.section) for s in segments] == [
        (1, 1, "1. DEFINITIONS"),
        (2, 2, "5 Payment Terms"),
    ]


def test_inline_heading_sets_section(tmp_path):
    pdf = make_pdf(
        tmp_path / "inline.pdf",
        [
            [
                "2. TERM",
                "2.2 Renewal. This Agreement shall automatically renew for one year.",
            ]
        ],
    )
    segments, _ = parse_document(str(pdf), PDF_MIME)
    assert len(segments) == 1
    assert segments[0].section == "2.2 Renewal"
    assert segments[0].text.startswith("2.2 Renewal. This Agreement")


def test_hyphenated_line_break_rejoined(tmp_path):
    pdf = make_pdf(
        tmp_path / "hyphen.pdf",
        [["Either party may give termi-\nnation notice in writing."]],
    )
    segments, _ = parse_document(str(pdf), PDF_MIME)
    assert segments[0].text == "Either party may give termination notice in writing."


def test_curly_quotes_normalized(tmp_path):
    # Base-14 PDF fonts cannot encode curly quotes, so use DOCX for this check.
    path = make_docx(
        tmp_path / "quotes.docx",
        [
            (
                "Normal",
                "Notice to the \u201cProvider\u201d at the Customer\u2019s option \u2014 in writing.",
            )
        ],
    )
    segments, _ = parse_document(str(path), DOCX_MIME)
    assert segments[0].text == 'Notice to the "Provider" at the Customer\'s option - in writing.'


def test_docx_heading_style_sets_section_and_no_page(tmp_path):
    path = make_docx(
        tmp_path / "doc.docx",
        [
            ("Title", "Master Agreement"),
            ("Normal", "This agreement is made between the parties named below."),
            ("Heading 2", "Confidentiality"),
            ("Normal", "Each party shall keep the other party's information secret."),
            ("Normal", ""),
        ],
    )
    segments, page_count = parse_document(str(path), DOCX_MIME)
    assert page_count is None
    assert [(s.seq, s.page, s.section) for s in segments] == [
        (1, None, "Master Agreement"),
        (2, None, "Confidentiality"),
    ]


def test_docx_table_rows_joined(tmp_path):
    document = docx.Document()
    document.add_paragraph("Fee schedule for the services provided under this agreement.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Support"
    table.rows[0].cells[1].text = "USD 1,000 per month"
    path = tmp_path / "table.docx"
    document.save(path)
    segments, _ = parse_document(str(path), DOCX_MIME)
    assert segments[-1].text == "Support | USD 1,000 per month"


def test_long_paragraph_split_at_sentences():
    text = " ".join(f"Sentence number {i} is here." for i in range(200))
    chunks = split_long(text)
    assert len(chunks) > 1
    assert all(len(c) <= 1500 for c in chunks)
    assert all(c.endswith(".") for c in chunks)
    assert " ".join(chunks) == text


def test_blank_pdf_raises_no_extractable_text(tmp_path):
    pdf = make_pdf(tmp_path / "blank.pdf", [[]])
    with pytest.raises(NoExtractableText):
        parse_document(str(pdf), PDF_MIME)


def test_garbage_pdf_raises_parse_error(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF-1.4 this is not really a pdf \x00\x01\x02")
    with pytest.raises(DocumentParseError):
        parse_document(str(path), PDF_MIME)


def test_encrypted_pdf_raises_parse_error(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Secret contract text that is long enough.")
    path = tmp_path / "enc.pdf"
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o")
    doc.close()
    with pytest.raises(DocumentParseError):
        parse_document(str(path), PDF_MIME)


def test_detect_mime_type(tmp_path):
    pdf = make_pdf(tmp_path / "a.pdf", [["Some text"]])
    path = make_docx(tmp_path / "a.docx", [("Normal", "Some text")])
    assert detect_mime_type("a.pdf", pdf.read_bytes()[:8]) == PDF_MIME
    assert detect_mime_type("A.PDF", pdf.read_bytes()[:8]) == PDF_MIME
    assert detect_mime_type("a.docx", path.read_bytes()[:8]) == DOCX_MIME
    assert detect_mime_type("setup.pdf", b"MZ\x90\x00\x03\x00\x00\x00") is None
    assert detect_mime_type("a.docx", pdf.read_bytes()[:8]) is None
    assert detect_mime_type("notes.txt", b"plain text") is None


def test_acme_txt_renewal_segment():
    segments, page_count = parse_document(str(ACME_TXT), TEXT_MIME)
    assert page_count is None
    assert segments[0].section == "SERVICES AGREEMENT"
    renewal = by_text(segments, RENEWAL_QUOTE)
    assert renewal.section == "2.2 Renewal"
    sections = {s.section for s in segments}
    assert {
        "2.1 Term",
        "2.2 Renewal",
        "4.2 Reporting Obligations",
        "6.1 Insurance",
    } <= sections
    assert [s.seq for s in segments] == list(range(1, len(segments) + 1))


def test_acme_pdf_matches_txt_with_pages():
    txt_segments, _ = parse_document(str(ACME_TXT), TEXT_MIME)
    pdf_segments, page_count = parse_document(str(ACME_PDF), PDF_MIME)
    assert page_count >= 2
    assert [(s.section, s.text) for s in pdf_segments] == [(s.section, s.text) for s in txt_segments]
    assert by_text(pdf_segments, RENEWAL_QUOTE).page == 1
    assert by_text(pdf_segments, "certificate of insurance").page == 2


def test_acme_docx_matches_txt():
    txt_segments, _ = parse_document(str(ACME_TXT), TEXT_MIME)
    docx_segments, page_count = parse_document(str(ACME_DOCX), DOCX_MIME)
    assert page_count is None
    assert [(s.section, s.text) for s in docx_segments] == [(s.section, s.text) for s in txt_segments]


@pytest.mark.parametrize("stem", ["conflicting/globex-hosting-agreement"])
@pytest.mark.parametrize("ext,mime", [(".pdf", PDF_MIME), (".docx", DOCX_MIME)])
def test_globex_samples_match_txt(stem, ext, mime):
    base = ROOT / "contracts/samples" / stem
    expected, _ = parse_document(str(base.with_suffix(".txt")), TEXT_MIME)
    segments, _ = parse_document(str(base.with_suffix(ext)), mime)
    assert [(s.section, s.text) for s in segments] == [(s.section, s.text) for s in expected]


def test_short_numbered_clause_is_not_dropped_as_heading(tmp_path):
    """Regression: '4.1 Client shall provide ...' (short, numbered) must stay a segment."""
    txt = tmp_path / "c.txt"
    txt.write_text(
        "4. Cooperation\n\n4.1 Client shall provide timely access to relevant personnel and information.\n\n"
        "5.1 Fees are payable within the period agreed with the relevant department.\n",
        encoding="utf-8",
    )
    segments, _ = parse_document(str(txt), "text/plain")
    texts = [s.text for s in segments]
    assert any("timely access" in t for t in texts)
    assert any("Fees are payable" in t for t in texts)
    assert segments[0].section == "4. Cooperation"
