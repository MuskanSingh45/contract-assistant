"""Render the .txt demo contracts into PDF and DOCX for the live upload demo.

Run from the repo root: .venv/bin/python tests/backend/documents/make_samples.py
"""

import sys
from pathlib import Path

import docx
import pymupdf

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))  # allow running as a plain script from anywhere

from backend.documents.normalizer import inline_heading, is_heading

SAMPLES = [
    # (txt path, heading to start a new page at, so the demo spans 2 pages)
    (ROOT / "contracts/samples/simple/acme-services-agreement.txt", "4. REPORTING"),
    (ROOT / "contracts/samples/conflicting/globex-hosting-agreement.txt", None),
    (ROOT / "contracts/samples/renewal/northstar-software-subscription.txt", None),
    (ROOT / "contracts/samples/ambiguous/meridian-consulting-agreement.txt", None),
    (ROOT / "contracts/samples/multiple-obligations/harbor-logistics-agreement.txt", "3. CUSTOMER OBLIGATIONS"),
]
A4 = pymupdf.paper_rect("a4")
MARGIN = 72
FONT_SIZE = 11
PARAGRAPH_GAP = 14


def read_paragraphs(txt: Path) -> list[tuple[str, bool]]:
    """(text, is_heading) per blank-line-separated paragraph."""
    paragraphs = [" ".join(p.split()) for p in txt.read_text(encoding="utf-8").split("\n\n")]
    return [(p, is_heading(p) and inline_heading(p) is None) for p in paragraphs if p]


def write_pdf(paragraphs: list[tuple[str, bool]], out: Path, page_break_before: str | None) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=A4.width, height=A4.height)
    y = MARGIN
    for text, heading in paragraphs:
        if text == page_break_before:
            page, y = doc.new_page(width=A4.width, height=A4.height), MARGIN
        font = "hebo" if heading else "helv"
        while True:
            rect = pymupdf.Rect(MARGIN, y, A4.width - MARGIN, A4.height - MARGIN)
            spare = page.insert_textbox(rect, text, fontsize=FONT_SIZE, fontname=font)
            if spare >= 0:
                break
            page, y = doc.new_page(width=A4.width, height=A4.height), MARGIN
        y = rect.y1 - spare + PARAGRAPH_GAP
    doc.save(out)
    doc.close()


def write_docx(paragraphs: list[tuple[str, bool]], out: Path) -> None:
    document = docx.Document()
    for i, (text, heading) in enumerate(paragraphs):
        if heading:
            document.add_heading(text, level=1 if i == 0 else 2)
        else:
            document.add_paragraph(text)
    document.save(out)


def main() -> None:
    for txt, page_break_before in SAMPLES:
        paragraphs = read_paragraphs(txt)
        write_pdf(paragraphs, txt.with_suffix(".pdf"), page_break_before)
        write_docx(paragraphs, txt.with_suffix(".docx"))


if __name__ == "__main__":
    main()
